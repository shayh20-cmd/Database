"""Seed data/spec_library.json from a ג2 DOCX, using the source's multilevel
list levels to reconstruct the sub-chapter / clause / sub-clause hierarchy.

Usage: python tools/spec-seed/seed.py "<path-to.docx>" [--preset-name "מבנה ציבור"]
Requires: python-docx."""
import argparse, json, os, sys, time
import docx
sys.path.insert(0, os.path.dirname(__file__))
from parser import (chapter_start, is_chapter_end, is_appendix, is_chapter_style,
                    detect_spine, build_chapter, new_id)

ROOT = os.path.join(os.path.dirname(__file__), '..', '..')
REF_PATH = os.path.join(ROOT, 'data', 'spec_chapter_reference.json')
OUT_PATH = os.path.join(ROOT, 'data', 'spec_library.json')


def load_reference():
    with open(REF_PATH, encoding='utf-8') as f:
        return {r['num']: r for r in json.load(f)}


def para_is_bold(p):
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs)


def para_numinfo(p):
    pPr = p._p.pPr
    if pPr is not None and pPr.numPr is not None:
        npr = pPr.numPr
        numId = npr.numId.val if npr.numId is not None else None
        ilvl = npr.ilvl.val if npr.ilvl is not None else 0
        return numId, ilvl
    return None, None


def read_paragraphs(path):
    """Extract paragraph dicts, skipping the table-of-contents region."""
    d = docx.Document(path)
    last_toc = -1
    for idx, p in enumerate(d.paragraphs):
        if p.style.name.startswith('toc'):
            last_toc = idx
    out = []
    for p in d.paragraphs[last_toc + 1:]:
        t = p.text.strip()
        if not t:
            continue
        numId, ilvl = para_numinfo(p)
        out.append({'style': p.style.name, 'text': t, 'bold': para_is_bold(p),
                    'numId': numId, 'ilvl': ilvl})
    return out


def group_chapters(paragraphs):
    """Split the flat paragraph stream into per-chapter item lists.
    Stops at the appendix ('נספחים'); closes a chapter at 'סוף פרק'."""
    chapters = []
    cur = None
    for it in paragraphs:
        text = it['text']
        if is_appendix(text) and it['style'].startswith('Heading'):
            break
        cs = chapter_start(text) if is_chapter_style(it['style']) else None
        if cs:
            num, name = cs
            cur = {'num': num, 'name': name, 'items': []}
            chapters.append(cur)
            continue
        if is_chapter_end(text):
            cur = None
            continue
        if cur is not None:
            cur['items'].append(it)
    return chapters


def parse_library(paragraphs, ref):
    chapters = []
    for ch in group_chapters(paragraphs):
        meta = ref.get(ch['num'], {'discipline': 'OTHER'})
        spine = detect_spine(ch['items'])
        sub_chapters = build_chapter(ch['items'], spine)
        chapters.append({'num': ch['num'],
                         'name': ref.get(ch['num'], {}).get('name', ch['name']),
                         'discipline': meta['discipline'],
                         'subChapters': sub_chapters})
    return chapters


def build_preset(chapters, name):
    return {'id': new_id(), 'name': name, 'buildingType': name,
            'selections': [{'chapterNum': c['num'],
                            'subChapterIds': [s['id'] for s in c['subChapters']]}
                           for c in chapters]}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('docx_path')
    ap.add_argument('--preset-name', default='מבנה ציבור')
    args = ap.parse_args()
    ref = load_reference()
    paragraphs = read_paragraphs(args.docx_path)
    chapters = parse_library(paragraphs, ref)
    library = {'_ts': int(time.time() * 1000), 'chapters': chapters,
               'presets': [build_preset(chapters, args.preset_name)]}
    with open(OUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(library, f, ensure_ascii=False, indent=1)
    n_sub = sum(len(c['subChapters']) for c in chapters)
    print(f'Wrote {OUT_PATH}: {len(chapters)} chapters, {n_sub} sub-chapters.')


if __name__ == '__main__':
    main()
